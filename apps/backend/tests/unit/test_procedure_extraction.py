from __future__ import annotations

import io

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from silo.services.procedure_extraction import ProcedureExtractionError, extract_document


def _pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(72, 720, "Primeira pagina do procedimento")
    document.showPage()
    document.drawString(72, 720, "Segunda pagina do procedimento")
    document.save()
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Procedimento DOCX extraido.")
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    buffer = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Slide PPTX extraido."
    presentation.save(buffer)
    return buffer.getvalue()


def test_extracts_pdf_text_with_page_locations() -> None:
    result = extract_document("guide.pdf", "application/pdf", _pdf_bytes())

    assert [chunk.page_number for chunk in result.chunks] == [1, 2]
    assert "Primeira pagina" in result.chunks[0].content
    assert "Segunda pagina" in result.chunks[1].content


@pytest.mark.parametrize(
    ("filename", "mime_type", "content", "expected"),
    [
        ("guide.txt", "text/plain", b"Texto simples do procedimento", "Texto simples"),
        ("guide.md", "text/markdown", b"# Procedimento\n\nConteudo Markdown", "Conteudo Markdown"),
    ],
)
def test_extracts_plain_text_formats(
    filename: str, mime_type: str, content: bytes, expected: str
) -> None:
    result = extract_document(filename, mime_type, content)

    assert len(result.chunks) == 1
    assert expected in result.chunks[0].content


def test_extracts_docx_paragraphs() -> None:
    result = extract_document(
        "guide.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        _docx_bytes(),
    )

    assert "Procedimento DOCX extraido" in result.chunks[0].content


def test_extracts_pptx_text_with_slide_location() -> None:
    result = extract_document(
        "guide.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        _pptx_bytes(),
    )

    assert result.chunks[0].slide_number == 1
    assert "Slide PPTX extraido" in result.chunks[0].content


def test_rejects_unsupported_and_empty_documents() -> None:
    with pytest.raises(ProcedureExtractionError, match="Formato de arquivo não suportado"):
        extract_document("guide.exe", "application/octet-stream", b"binary")

    with pytest.raises(ProcedureExtractionError, match="não contém texto"):
        extract_document("guide.txt", "text/plain", b"  \n")


def test_rejects_mime_and_extension_mismatch() -> None:
    with pytest.raises(ProcedureExtractionError, match="MIME incompatível"):
        extract_document("guide.pdf", "text/plain", _pdf_bytes())

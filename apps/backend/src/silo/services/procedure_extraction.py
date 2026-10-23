from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from silo.ai.chunking import estimate_tokens


class ProcedureExtractionError(ValueError):
    """Raised when an attachment cannot be safely converted into searchable text."""


@dataclass(frozen=True, slots=True)
class ExtractedChunk:
    content: str
    index: int
    token_count: int
    page_number: int | None = None
    slide_number: int | None = None


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    chunks: tuple[ExtractedChunk, ...]


_MIME_BY_SUFFIX: dict[str, frozenset[str]] = {
    ".pdf": frozenset({"application/pdf"}),
    ".txt": frozenset({"text/plain"}),
    ".md": frozenset({"text/markdown", "text/plain"}),
    ".docx": frozenset(
        {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    ),
    ".pptx": frozenset(
        {"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
    ),
}


def extract_document(filename: str, mime_type: str, content: bytes) -> ExtractedDocument:
    suffix = PurePath(filename).suffix.lower()
    accepted_mimes = _MIME_BY_SUFFIX.get(suffix)
    if accepted_mimes is None:
        raise ProcedureExtractionError("Formato de arquivo não suportado.")
    if mime_type not in accepted_mimes:
        raise ProcedureExtractionError("MIME incompatível com a extensão do arquivo.")
    if not content:
        raise ProcedureExtractionError("O arquivo não contém texto pesquisável.")

    try:
        if suffix == ".pdf":
            chunks = _extract_pdf(content)
        elif suffix in {".txt", ".md"}:
            chunks = _plain_text_chunks(content)
        elif suffix == ".docx":
            chunks = _extract_docx(content)
        else:
            chunks = _extract_pptx(content)
    except ProcedureExtractionError:
        raise
    except Exception as exc:
        raise ProcedureExtractionError("Não foi possível extrair o conteúdo do arquivo.") from exc

    if not chunks:
        raise ProcedureExtractionError("O arquivo não contém texto pesquisável.")
    return ExtractedDocument(chunks=tuple(chunks))


def _extract_pdf(content: bytes) -> list[ExtractedChunk]:
    if not content.startswith(b"%PDF-"):
        raise ProcedureExtractionError("Arquivo PDF inválido.")
    reader = PdfReader(io.BytesIO(content))
    chunks: list[ExtractedChunk] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            chunks.append(
                ExtractedChunk(
                    content=text,
                    index=len(chunks),
                    token_count=estimate_tokens(text),
                    page_number=page_index,
                )
            )
    return chunks


def _plain_text_chunks(content: bytes) -> list[ExtractedChunk]:
    try:
        text = content.decode("utf-8").strip()
    except UnicodeDecodeError as exc:
        raise ProcedureExtractionError("Arquivo de texto não está em UTF-8.") from exc
    if not text:
        return []
    return [ExtractedChunk(content=text, index=0, token_count=estimate_tokens(text))]


def _extract_docx(content: bytes) -> list[ExtractedChunk]:
    document = Document(io.BytesIO(content))
    text = "\n\n".join(
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    )
    if not text:
        return []
    return [ExtractedChunk(content=text, index=0, token_count=estimate_tokens(text))]


def _extract_pptx(content: bytes) -> list[ExtractedChunk]:
    presentation = Presentation(io.BytesIO(content))
    chunks: list[ExtractedChunk] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text")]
        text = "\n".join(value for value in texts if value)
        if text:
            chunks.append(
                ExtractedChunk(
                    content=text,
                    index=len(chunks),
                    token_count=estimate_tokens(text),
                    slide_number=slide_index,
                )
            )
    return chunks
